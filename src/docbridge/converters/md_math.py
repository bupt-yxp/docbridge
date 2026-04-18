from __future__ import annotations

import base64
import html as html_module
import logging
import re
from io import BytesIO

from bs4 import BeautifulSoup, NavigableString, Tag

logger = logging.getLogger(__name__)

_DOC_INLINE_CLS = "docbridge-math docbridge-math-inline"
_DOC_DISPLAY_CLS = "docbridge-math docbridge-math-display"

# 与 _replace_multiline_single_dollar_display 共用：行首/行尾单独一行的 $
_MULTILINE_SINGLE_DOLLAR_DISPLAY = re.compile(
    r"(?m)^[ \t]*\$[ \t]*\n([\s\S]*?)\n[ \t]*\$[ \t]*(?:\n|$)",
)


def _is_dollar_escaped(s: str, dollar_index: int) -> bool:
    """True if the '$' at dollar_index is preceded by an odd number of backslashes."""
    k = dollar_index - 1
    n = 0
    while k >= 0 and s[k] == "\\":
        n += 1
        k -= 1
    return n % 2 == 1


def substitute_tex_delimiters(md_text: str) -> str:
    """
    Replace $...$, $$...$$, \\(...\\), \\[...\\] with raw HTML placeholders so that
    (1) underscores inside math are not parsed as Markdown emphasis, and
    (2) downstream can emit OMML (Word) or SVG (PDF).
    """
    text = md_text
    text = _replace_display_brackets(text)
    text = _replace_inline_parens(text)
    text = _replace_dollar_delims(text)
    return text


def _replace_multiline_single_dollar_display(text: str) -> str:
    """
    LaTeX-style display: opening `$` on its own line, body, closing `$` on its own line.
    Must run after `$$...$$` is stripped so we do not split `$$`.
    """

    def repl(m: re.Match[str]) -> str:
        return _wrap_display(m.group(1).strip())

    return _MULTILINE_SINGLE_DOLLAR_DISPLAY.sub(repl, text)


def normalize_multiline_dollar_display_for_pandoc(md_text: str) -> str:
    """
    将「单独成行」的 `$ … $` 显示块规范化为 `$$ … $$`。

    Pandoc 对前者解析不稳定，易导致 LaTeX 把 `f'(x)` 等留在正文而出
    ``! Missing $ inserted``。统一为 `$$` 后由 tex_math_dollars 走正规 display math。
    """
    def repl(m: re.Match[str]) -> str:
        inner = m.group(1).strip()
        return f"\n$$\n{inner}\n$$\n"

    return _MULTILINE_SINGLE_DOLLAR_DISPLAY.sub(repl, md_text)


def _b64_encode(s: str) -> str:
    return base64.standard_b64encode(s.encode("utf-8")).decode("ascii")


def _wrap_display(latex: str) -> str:
    return f'\n\n<div class="{_DOC_DISPLAY_CLS}" data-latex="{_b64_encode(latex)}"></div>\n\n'


def _wrap_inline(latex: str) -> str:
    return f'<span class="{_DOC_INLINE_CLS}" data-latex="{_b64_encode(latex)}"></span>'


def _replace_display_brackets(text: str) -> str:
    pattern = re.compile(
        r"\\\[([\s\S]*?)\\\]",
        re.MULTILINE,
    )

    def repl(m: re.Match[str]) -> str:
        return _wrap_display(m.group(1).strip())

    return pattern.sub(repl, text)


def _replace_inline_parens(text: str) -> str:
    out: list[str] = []
    i, n = 0, len(text)
    while i < n:
        if i + 1 < n and text[i : i + 2] == "\\(":
            j = i + 2
            depth = 1
            while j + 1 < n:
                if text[j : j + 2] == "\\(":
                    depth += 1
                    j += 2
                    continue
                if text[j : j + 2] == "\\)":
                    depth -= 1
                    if depth == 0:
                        latex = text[i + 2 : j].strip()
                        out.append(_wrap_inline(latex))
                        i = j + 2
                        break
                    j += 2
                    continue
                j += 1
            else:
                out.append(text[i])
                i += 1
            continue
        out.append(text[i])
        i += 1
    return "".join(out)


def _replace_double_dollar(text: str) -> str:
    """Replace $$...$$ (may span lines) with display placeholders."""
    out: list[str] = []
    i, n = 0, len(text)
    while i < n:
        if i + 1 < n and text[i : i + 2] == "$$" and not _is_dollar_escaped(text, i):
            j = i + 2
            found = False
            while j + 1 < n:
                if text[j : j + 2] == "$$" and not _is_dollar_escaped(text, j):
                    latex = text[i + 2 : j].strip()
                    out.append(_wrap_display(latex))
                    i = j + 2
                    found = True
                    break
                j += 1
            if not found:
                out.append(text[i])
                i += 1
            continue
        out.append(text[i])
        i += 1
    return "".join(out)


def _replace_inline_single_dollar(text: str) -> str:
    """Replace single-line $...$ with inline placeholders (no newline inside)."""
    out: list[str] = []
    i, n = 0, len(text)
    while i < n:
        if i + 1 < n and text[i : i + 2] == "$$":
            out.append(text[i])
            i += 1
            continue

        if text[i] == "$" and not _is_dollar_escaped(text, i):
            j = i + 1
            while j < n:
                if text[j] == "$" and not _is_dollar_escaped(text, j):
                    latex = text[i + 1 : j]
                    if "\n" in latex:
                        out.append(text[i])
                        i += 1
                        break
                    out.append(_wrap_inline(latex.strip()))
                    i = j + 1
                    break
                j += 1
            else:
                out.append(text[i])
                i += 1
            continue

        out.append(text[i])
        i += 1
    return "".join(out)


def _replace_dollar_delims(text: str) -> str:
    """$$...$$ → $\n...\n$ (display) → single-line $...$ (inline)."""
    text = _replace_double_dollar(text)
    text = _replace_multiline_single_dollar_display(text)
    text = _replace_inline_single_dollar(text)
    return text


def decode_latex_data_attr(b64: str) -> str:
    return base64.standard_b64decode(b64.encode("ascii")).decode("utf-8")


def _repair_omml_xml_for_word(xml: str) -> str:
    """
    mathml2omml 对部分记号（典型为 \\bar 的「上划线」）会生成 **非良构** 的 ``m:groupChr``
   （提前闭合标签），python-docx 的 ``parse_xml`` 严格校验失败，上游只能回退为 ``$...$`` 纯文字。

    使用 lxml 的 recover 解析器纠错后再序列化，可得 Word 可接受的 OMML。
    """
    from lxml import etree

    parser_strict = etree.XMLParser(remove_blank_text=False, resolve_entities=False)
    try:
        etree.fromstring(xml.encode("utf-8"), parser_strict)
        return xml
    except etree.XMLSyntaxError:
        pass

    parser_recover = etree.XMLParser(recover=True, remove_blank_text=False, resolve_entities=False)
    wrapped = (
        '<docbridge xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        + xml
        + "</docbridge>"
    )
    root = etree.fromstring(wrapped.encode("utf-8"), parser_recover)
    return etree.tostring(root[0], encoding="unicode")


def latex_to_omml_elements(latex: str, *, display: bool) -> str:
    """Return XML string: either oMathPara (display) or w:r wrapping m:oMath (inline)."""
    import latex2mathml.converter as lc
    import mathml2omml

    mml = lc.convert(latex, display="block" if display else "inline")
    omml = mathml2omml.convert(mml)
    m_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    if display:
        raw = f'<m:oMathPara xmlns:m="{m_ns}">' f"{omml}" "</m:oMathPara>"
        return _repair_omml_xml_for_word(raw)
    raw = f'<w:r xmlns:w="{w_ns}" xmlns:m="{m_ns}">' f"{omml}" "</w:r>"
    return _repair_omml_xml_for_word(raw)


def append_omml_to_paragraph(paragraph, latex: str, *, display: bool) -> None:
    from docx.oxml import parse_xml

    xml = latex_to_omml_elements(latex, display=display)
    paragraph._p.append(parse_xml(xml))


def expand_math_tags_for_pdf(html_fragment: str) -> str:
    """Replace docbridge math placeholders with matplotlib-rendered SVG for WeasyPrint."""
    wrapped = f"<html><body>{html_fragment}</body></html>"
    soup = BeautifulSoup(wrapped, "html.parser")
    body = soup.body
    if not body:
        return html_fragment

    for el in body.find_all(class_=lambda c: bool(c) and "docbridge-math" in c):
        if not isinstance(el, Tag):
            continue
        tag = el
        b64 = tag.get("data-latex")
        if not b64:
            continue
        try:
            latex = decode_latex_data_attr(b64)
        except Exception:
            tag.replace_with(f"[invalid math: {b64[:20]}…]")
            continue
        display = "docbridge-math-display" in (tag.get("class") or [])
        svg_html = _latex_to_svg_html(latex, display=display)
        if svg_html is None:
            fb = soup.new_tag("span", attrs={"class": "docbridge-math-fallback"})
            fb.append(NavigableString(html_module.escape(latex)))
            tag.replace_with(fb)
            continue
        frag = BeautifulSoup(svg_html, "html.parser")
        tag.replace_with(frag)

    inner = "".join(str(c) for c in body.contents)
    return inner


def _normalize_latex_for_matplotlib(s: str) -> str:
    """Map common AMS/LaTeX to matplotlib mathtext subset (still lossy for some envs)."""
    t = s.strip()
    if not t:
        return t
    t = t.replace("\\dfrac", "\\frac")
    t = re.sub(r"\\lim\s*\\limits", r"\\lim", t)
    t = re.sub(r"\\iint", r"\\int\\!\\!\\int", t)
    # \begin{pmatrix} / cases / matrix：当前 matplotlib mathtext 不支持 amsmath 环境，保留原串；
    # 渲染失败时在 PDF 中降级为等宽文本（见 expand_math_tags_for_pdf）。
    return t


def _latex_to_svg_html(latex: str, *, display: bool) -> str | None:
    try:
        import matplotlib

        matplotlib.use("Agg")
    except ImportError:
        return None

    normalized = _normalize_latex_for_matplotlib(latex)
    return _matplotlib_latex_to_svg_string(normalized, display=display)


def _matplotlib_latex_to_svg_string(latex: str, *, display: bool) -> str | None:
    import matplotlib.pyplot as plt

    fontsize = 13 if display else 11
    pad = 0.1 if display else 0.015
    w_in = min(6.4, 0.35 + 0.022 * max(len(latex), 8))
    h_in = 0.52 if display else 0.32
    try:
        fig = plt.figure(figsize=(w_in, h_in))
        ax = fig.add_axes([0, 0, 1, 1])
        ax.axis("off")
        if display:
            ax.text(
                0.5,
                0.5,
                f"${latex}$",
                fontsize=fontsize,
                ha="center",
                va="center",
                transform=ax.transAxes,
            )
        else:
            ax.text(
                0.0,
                0.5,
                f"${latex}$",
                fontsize=fontsize,
                ha="left",
                va="center",
                transform=ax.transAxes,
            )
        buf = BytesIO()
        fig.savefig(
            buf,
            format="svg",
            bbox_inches="tight",
            pad_inches=pad,
            transparent=True,
        )
        plt.close(fig)
        svg = buf.getvalue().decode("utf-8")
    except Exception as e:
        logger.debug("matplotlib mathtext render failed: %s", e)
        return None

    if display:
        return f'<div class="docbridge-math-svg docbridge-math-svg-display">{svg}</div>'
    return f'<span class="docbridge-math-svg docbridge-math-svg-inline">{svg}</span>'
